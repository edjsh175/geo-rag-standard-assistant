"""Structure-preserving document parsing for uploaded-document indexing."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Any

from app.services.document_text_extractor import DocumentTextExtractor


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    markdown: str
    metadata: dict[str, Any] = field(default_factory=dict)


class DocumentParser:
    """Normalize supported documents into Markdown-like structured text."""

    def __init__(self, extractor: DocumentTextExtractor | Any | None = None) -> None:
        self.extractor = extractor or DocumentTextExtractor()

    def parse(self, path: Path, content_type: str | None = None) -> ParsedDocument:
        suffix = path.suffix.lower()
        normalized_content_type = (content_type or "").split(";", 1)[0].strip().lower()

        if suffix == ".md" or normalized_content_type == "text/markdown":
            return ParsedDocument(markdown=self._normalize_markdown(path.read_text(encoding="utf-8")))
        if suffix == ".docx" or normalized_content_type.endswith("wordprocessingml.document"):
            return ParsedDocument(markdown=self._read_docx(path))
        if suffix in {".xlsx", ".xlsm"} or normalized_content_type.endswith("spreadsheetml.sheet"):
            return ParsedDocument(markdown=self._read_xlsx(path))

        extracted = self.extractor.extract(path, content_type)
        return ParsedDocument(
            markdown=self._normalize_markdown(extracted.text),
            metadata=dict(extracted.metadata or {}),
        )

    @classmethod
    def _normalize_markdown(cls, text: str) -> str:
        normalized_lines: list[str] = []
        in_fence = False
        for raw_line in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            if raw_line.lstrip().startswith("```"):
                in_fence = not in_fence
                normalized_lines.append(raw_line.rstrip())
                continue
            if in_fence:
                normalized_lines.append(raw_line.rstrip())
                continue

            line = raw_line.replace("\u00a0", " ").rstrip()
            line = re.sub(r"[ \t]+", " ", line).strip()
            normalized_lines.append(line)

        collapsed: list[str] = []
        previous_blank = False
        for line in normalized_lines:
            is_blank = not line
            if is_blank and previous_blank:
                continue
            collapsed.append(line)
            previous_blank = is_blank
        return "\n".join(collapsed).strip()

    @classmethod
    def _read_docx(cls, path: Path) -> str:
        from docx import Document as DocxDocument

        document = DocxDocument(str(path))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = cls._normalize_inline(paragraph.text)
            if not text:
                continue
            style_name = (paragraph.style.name or "").strip().lower() if paragraph.style else ""
            heading_match = re.match(r"heading\s*(\d+)", style_name)
            if heading_match:
                level = max(1, min(6, int(heading_match.group(1))))
                blocks.append(f"{'#' * level} {text}")
            else:
                blocks.append(text)

        for table in document.tables:
            rows = [[cls._normalize_inline(cell.text) for cell in row.cells] for row in table.rows]
            table_text = cls._markdown_table(rows)
            if table_text:
                blocks.append(table_text)
        return cls._normalize_markdown("\n\n".join(blocks))

    @classmethod
    def _read_xlsx(cls, path: Path) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        blocks: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [cls._cell_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append(values)
            if not rows:
                continue
            blocks.append(f"# {sheet.title}")
            blocks.append(cls._markdown_table(rows))
        return cls._normalize_markdown("\n\n".join(blocks))

    @staticmethod
    def _normalize_inline(value: str) -> str:
        return re.sub(r"\s+", " ", (value or "").replace("\u00a0", " ")).strip()

    @classmethod
    def _markdown_table(cls, rows: list[list[str]]) -> str:
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        header = normalized[0]
        lines = [
            "| " + " | ".join(cls._escape_table_cell(value) for value in header) + " |",
            "| " + " | ".join("---" for _ in range(width)) + " |",
        ]
        lines.extend(
            "| " + " | ".join(cls._escape_table_cell(value) for value in row) + " |"
            for row in normalized[1:]
        )
        return "\n".join(lines)

    @staticmethod
    def _escape_table_cell(value: str) -> str:
        return str(value or "").replace("|", "\\|").strip()

    @staticmethod
    def _cell_text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
