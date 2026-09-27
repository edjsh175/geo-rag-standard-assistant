"""Comprehensive tests for Phase M: Document Ingestion, Parsing, Chunking, Lineage & Health Audit."""

from __future__ import annotations

from contextlib import asynccontextmanager
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
import pytest
from pydantic import ValidationError

from app.core.config import Settings, validate_embedding_dimension_invariants
from app.services.document_chunker import (
    DocumentChunk,
    DocumentChunker,
    compute_chunk_uid,
    compute_content_hash,
)
from app.services.document_index_audit_service import (
    AuditIssue,
    CorpusAuditReport,
    DocumentIndexAuditService,
)
from app.services.document_indexing_service import DocumentIndexingService
from app.services.document_parser import DocumentParser, ParsedBlock, ParsedDocument
from app.services.document_text_extractor import DocumentTextExtractor, ExtractedDocumentText


# ==============================================================================
# M-01: DOCX Paragraph & Table Order Preservation
# ==============================================================================
def test_m01_docx_ordered_traversal_preserves_reading_flow(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "mixed_flow.docx"
    doc = Document()
    doc.add_heading("总规划说明", level=1)
    doc.add_paragraph("第一节：背景介绍。")

    t1 = doc.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "指标"
    t1.cell(0, 1).text = "现状"
    t1.cell(1, 0).text = "用地面积"
    t1.cell(1, 1).text = "100公顷"

    doc.add_paragraph("第二节：核心管控要求。")

    t2 = doc.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "控制项"
    t2.cell(0, 1).text = "红线"
    t2.cell(1, 0).text = "容积率"
    t2.cell(1, 1).text = "2.5"

    doc.add_paragraph("第三节：实施保障机制。")
    doc.save(path)

    parsed = DocumentParser().parse(path)

    # 1. Assert blocks are populated and in exact body sequence
    kinds = [b.kind for b in parsed.blocks]
    assert kinds == ["heading", "paragraph", "table", "paragraph", "table", "paragraph"], (
        f"Reading flow was corrupted: got kinds {kinds}"
    )

    # 2. Assert markdown retains the exact interwoven sequence
    assert "总规划说明" in parsed.markdown
    idx_p1 = parsed.markdown.find("第一节：背景介绍。")
    idx_t1 = parsed.markdown.find("| 用地面积 | 100公顷 |")
    idx_p2 = parsed.markdown.find("第二节：核心管控要求。")
    idx_t2 = parsed.markdown.find("| 容积率 | 2.5 |")
    idx_p3 = parsed.markdown.find("第三节：实施保障机制。")

    assert 0 <= idx_p1 < idx_t1 < idx_p2 < idx_t2 < idx_p3, (
        "DOCX markdown serialization destroyed paragraph-table relative ordering!"
    )


# ==============================================================================
# M-02: PDF Page Provenance & Chunk Page Propagation
# ==============================================================================
def test_m02_pdf_page_provenance_propagates_to_chunks() -> None:
    # Construct a mock extractor returning multi-page text
    mock_extractor = MagicMock(spec=DocumentTextExtractor)
    mock_extractor.extract.return_value = ExtractedDocumentText(
        text="第1页前言说明。\n\n第2页指标核准条目。\n\n第3页附则规定。",
        metadata={"page_count": 3},
        pages=[
            {"page_number": 1, "text": "第1页前言说明。"},
            {"page_number": 2, "text": "第2页指标核准条目。"},
            {"page_number": 3, "text": "第3页附则规定。"},
        ],
    )

    parser = DocumentParser(extractor=mock_extractor)
    parsed = parser.parse(Path("mock.pdf"), "application/pdf")

    assert len(parsed.blocks) == 3
    assert [b.page_number for b in parsed.blocks] == [1, 2, 3]

    chunker = DocumentChunker(chunk_size=100, chunk_overlap=10)
    chunks = chunker.chunk(parsed)

    assert len(chunks) == 3
    assert chunks[0].page_number == 1
    assert "第1页前言说明" in chunks[0].content
    assert chunks[1].page_number == 2
    assert "第2页指标核准条目" in chunks[1].content
    assert chunks[2].page_number == 3
    assert "第3页附则规定" in chunks[2].content


# ==============================================================================
# M-03 / M-04: Stable Chunk Identity & Lineage Adjacency
# ==============================================================================
def test_m03_deterministic_chunk_uid_stability() -> None:
    doc_id = "d1111111-1111-1111-1111-111111111111"
    section = "总则 > 第一条"
    c_hash = compute_content_hash("本规划适用于城市开发边界内所有新建项目。")

    uid1 = compute_chunk_uid(doc_id, section, 0, c_hash)
    uid2 = compute_chunk_uid(doc_id, section, 0, c_hash)
    assert uid1 == uid2, "Same inputs must produce deterministic identical chunk_uid"

    # Changing index or content changes UID
    diff_index_uid = compute_chunk_uid(doc_id, section, 1, c_hash)
    assert diff_index_uid != uid1


@pytest.mark.asyncio
async def test_m04_chunk_lineage_and_adjacency_closure() -> None:
    doc_id = "d2222222-2222-2222-2222-222222222222"
    version_id = "v2222222-2222-2222-2222-222222222222"

    mock_repo = MagicMock()
    mock_repo.get_indexing_payload = AsyncMock(
        return_value={
            "document_id": doc_id,
            "version_id": version_id,
            "title": "测试标准规范",
            "filename": "standard.md",
            "metadata": {"domain": "urban_planning"},
        }
    )
    mock_repo.mark_job_running = AsyncMock()
    mock_repo.update_job_stage = AsyncMock()
    mock_repo.replace_chunks = AsyncMock()
    mock_repo.mark_job_succeeded = AsyncMock()

    mock_storage = MagicMock()
    mock_storage.download_version_to_temp = AsyncMock(return_value=Path("dummy.md"))

    mock_parser = MagicMock()
    mock_parser.parse.return_value = ParsedDocument(
        markdown="# 第一条\n\n内容A。\n\n# 第二条\n\n内容B。\n\n# 第三条\n\n内容C。",
        blocks=[
            ParsedBlock(kind="heading", content="# 第一条", order=1, header_path="第一条"),
            ParsedBlock(kind="paragraph", content="内容A。", order=2, header_path="第一条"),
            ParsedBlock(kind="heading", content="# 第二条", order=3, header_path="第二条"),
            ParsedBlock(kind="paragraph", content="内容B。", order=4, header_path="第二条"),
            ParsedBlock(kind="heading", content="# 第三条", order=5, header_path="第三条"),
            ParsedBlock(kind="paragraph", content="内容C。", order=6, header_path="第三条"),
        ],
    )

    mock_embed = MagicMock()
    mock_embed.embed_texts = AsyncMock(return_value=[[0.1] * 2048, [0.2] * 2048, [0.3] * 2048])

    service = DocumentIndexingService(
        repository=mock_repo,
        storage=mock_storage,
        parser=mock_parser,
        embedding_provider=mock_embed,
    )

    await service.run_job("job-123")

    mock_repo.replace_chunks.assert_awaited_once()
    saved_kwargs = mock_repo.replace_chunks.await_args.kwargs
    chunks = saved_kwargs["chunks"]
    assert len(chunks) == 3

    # Check Adjacency Lineage
    assert chunks[0]["metadata"]["prev_chunk_uid"] is None
    assert chunks[0]["metadata"]["next_chunk_uid"] == chunks[1]["chunk_uid"]

    assert chunks[1]["metadata"]["prev_chunk_uid"] == chunks[0]["chunk_uid"]
    assert chunks[1]["metadata"]["next_chunk_uid"] == chunks[2]["chunk_uid"]

    assert chunks[2]["metadata"]["prev_chunk_uid"] == chunks[1]["chunk_uid"]
    assert chunks[2]["metadata"]["next_chunk_uid"] is None

    # Check Chunk Policy ID and Content Role
    for c in chunks:
        assert c["metadata"]["chunk_policy_id"] == "section_based_v1"
        assert "content_role" in c["metadata"]
        assert "source_element_orders" in c["metadata"]


# ==============================================================================
# M-06: Atomic Table Row Grouping with Header Preservation
# ==============================================================================
def test_m06_large_table_row_grouping_preserves_headers() -> None:
    # Build a table that exceeds a 150-char chunk_size limit
    header = "| 项目序号 | 管控指标名称 | 规划指标数值 | 备注说明 |"
    sep = "| --- | --- | --- | --- |"
    rows = [f"| {i:02d} | 指标类型_{i} | 数值_{i * 10} | 依据规划第{i}条实施 |" for i in range(1, 15)]
    full_table = "\n".join([header, sep] + rows)

    parsed = ParsedDocument(
        markdown=full_table,
        blocks=[
            ParsedBlock(
                kind="table",
                content=full_table,
                order=1,
                header_path="指标附表",
            )
        ],
    )

    chunker = DocumentChunker(chunk_size=150, chunk_overlap=10)
    chunks = chunker.chunk(parsed)

    assert len(chunks) > 1, "Large table must be split into multiple row-group chunks"

    for chunk in chunks:
        # Every chunk must preserve the table header and separator
        assert "| 项目序号 | 管控指标名称 |" in chunk.content
        assert "| --- | --- |" in chunk.content
        assert chunk.chunk_policy_id == "table_rowgroup_v1"
        assert chunk.content_role == "table"
        assert chunk.header_path == "指标附表"

    # All data rows should appear across the chunks
    for i in range(1, 15):
        assert any(f"| {i:02d} |" in chunk.content for chunk in chunks), f"Row {i} was lost during table split!"


# ==============================================================================
# M-07: Schema Vector Dimension 2048 Invariant
# ==============================================================================
def test_m07_embedding_dimension_schema_invariant() -> None:
    # Valid config
    valid_cfg = Settings(PG_VECTOR_DIMENSION=2048, OLLAMA_EMBEDDING_DIMENSIONS=2048)
    validate_embedding_dimension_invariants(valid_cfg)

    # Invalid PG_VECTOR_DIMENSION raises
    with pytest.raises(ValidationError):
        Settings(PG_VECTOR_DIMENSION=1536)

    # Invalid OLLAMA_EMBEDDING_DIMENSIONS raises
    with pytest.raises(ValidationError):
        Settings(OLLAMA_EMBEDDING_DIMENSIONS=768)


# ==============================================================================
# M-08: DocumentIndexAuditService Integrity Gates
# ==============================================================================
@pytest.mark.asyncio
async def test_m08_corpus_audit_service_healthy() -> None:
    mock_session = AsyncMock()

    class FakeResult:
        def __init__(self, rows):
            self._rows = rows
        def mappings(self):
            return self
        def first(self):
            return self._rows[0] if self._rows else None
        def all(self):
            return self._rows

    # Mock count queries and empty issue queries
    def execute_side_effect(sql, *args, **kwargs):
        sql_str = str(sql).strip()
        if "total_docs" in sql_str:
            return FakeResult([{"total_docs": 5, "total_chunks": 42, "total_jobs": 10}])
        return FakeResult([])

    mock_session.execute = AsyncMock(side_effect=execute_side_effect)

    @asynccontextmanager
    async def mock_session_factory():
        yield mock_session

    audit_svc = DocumentIndexAuditService(session_factory=mock_session_factory)
    report = await audit_svc.audit_corpus()

    assert report.is_healthy is True
    assert report.total_documents == 5
    assert report.total_chunks == 42
    assert report.summary["error_count"] == 0
    assert len(report.issues) == 0


@pytest.mark.asyncio
async def test_m08_corpus_audit_service_detects_all_issue_categories() -> None:
    mock_session = AsyncMock()

    class FakeResult:
        def __init__(self, rows):
            self._rows = rows
        def mappings(self):
            return self
        def first(self):
            return self._rows[0] if self._rows else None
        def all(self):
            return self._rows

    def execute_side_effect(sql, *args, **kwargs):
        sql_str = str(sql).strip()
        if "total_docs" in sql_str:
            return FakeResult([{"total_docs": 10, "total_chunks": 50, "total_jobs": 20}])
        if "indexed_documents_without_chunks" in sql_str or "COUNT(c.id) = 0" in sql_str:
            return FakeResult([{"document_id": "doc-empty-1", "title": "Empty Doc", "index_status": "indexed"}])
        if "stale_version_chunks" in sql_str or "c.version_id != d.current_version_id" in sql_str:
            return FakeResult([{"chunk_id": "chunk-stale-1", "document_id": "doc-1", "chunk_version_id": "v1", "current_version_id": "v2"}])
        if "deleted_documents_with_chunks" in sql_str or "d.deleted_at IS NOT NULL" in sql_str:
            return FakeResult([{"chunk_id": "chunk-deleted-1", "document_id": "doc-del", "deleted_at": "2026-09-27"}])
        if "duplicate_chunk_uids" in sql_str or "HAVING COUNT(*) > 1" in sql_str:
            return FakeResult([{"chunk_uid": "uid-dup-1", "count": 2}])
        if "embedding_dimension_mismatch" in sql_str or "vector_dims(embedding) != 2048" in sql_str:
            return FakeResult([{"chunk_id": "chunk-bad-dim-1", "document_id": "doc-2", "anomaly_type": "dimension_mismatch"}])
        if "stuck_jobs" in sql_str or "status IN ('queued', 'running')" in sql_str:
            return FakeResult([{"job_id": "job-stuck-1", "document_id": "doc-3", "status": "running", "stage": "embedding", "updated_at": "2026-09-27"}])
        if "status_conflicts" in sql_str or "d.index_status = 'indexed' AND j.status = 'failed'" in sql_str:
            return FakeResult([{"document_id": "doc-conflict-1", "doc_status": "indexed", "job_status": "failed", "job_id": "job-f1"}])
        if "chunk_length_anomalies" in sql_str or "LENGTH(content) < 5" in sql_str:
            return FakeResult([{"chunk_id": "chunk-short-1", "document_id": "doc-4", "length": 3}])
        return FakeResult([])

    mock_session.execute = AsyncMock(side_effect=execute_side_effect)

    @asynccontextmanager
    async def mock_session_factory():
        yield mock_session

    audit_svc = DocumentIndexAuditService(session_factory=mock_session_factory)
    report = await audit_svc.audit_corpus()

    assert report.is_healthy is False
    assert report.summary["error_count"] > 0
    assert report.summary["warning_count"] > 0

    cats = set(report.issue_counts.keys())
    assert "indexed_documents_without_chunks" in cats
    assert "stale_version_chunks" in cats
    assert "deleted_documents_with_chunks" in cats
    assert "duplicate_chunk_uids" in cats
    assert "embedding_dimension_mismatch" in cats
    assert "stuck_jobs" in cats
    assert "status_conflicts" in cats
    assert "chunk_length_anomalies" in cats
