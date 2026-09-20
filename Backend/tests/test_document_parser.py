from __future__ import annotations

from app.services.document_parser import DocumentParser


def test_markdown_parser_preserves_headings_tables_and_fenced_code(tmp_path) -> None:
    path = tmp_path / "planning.md"
    path.write_text(
        """# 总则

适用范围。

| 编号 | 要求 |
| --- | --- |
| 1 | 保留表格 |

```python
def rule():
    return "保留代码块"
```
""",
        encoding="utf-8",
    )

    parsed = DocumentParser().parse(path, "text/markdown")

    assert parsed.markdown.startswith("# 总则")
    assert "| 编号 | 要求 |" in parsed.markdown
    assert "```python\ndef rule():\n    return \"保留代码块\"\n```" in parsed.markdown
    assert "适用范围。  \n" not in parsed.markdown


def test_docx_parser_keeps_heading_and_table_as_markdown(tmp_path) -> None:
    from docx import Document

    path = tmp_path / "standard.docx"
    document = Document()
    document.add_heading("规划要求", level=1)
    document.add_paragraph("  第一条\u00a0  建设边界应明确。  ")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "容积率"
    table.cell(1, 1).text = "2.0"
    document.save(path)

    parsed = DocumentParser().parse(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "# 规划要求" in parsed.markdown
    assert "第一条 建设边界应明确。" in parsed.markdown
    assert "| 指标 | 值 |" in parsed.markdown
    assert "| 容积率 | 2.0 |" in parsed.markdown


def test_xlsx_parser_converts_each_sheet_to_markdown_table(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "indicators.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "控制指标"
    sheet.append(["指标", "值"])
    sheet.append(["容积率", 2.0])
    second = workbook.create_sheet("备注")
    second.append(["说明"])
    second.append(["用于测试"])
    workbook.save(path)

    parsed = DocumentParser().parse(
        path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert "# 控制指标" in parsed.markdown
    assert "| 指标 | 值 |" in parsed.markdown
    assert "| 容积率 | 2 |" in parsed.markdown
    assert "# 备注" in parsed.markdown
    assert "| 说明 |" in parsed.markdown
    assert "| 用于测试 |" in parsed.markdown
